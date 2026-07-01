from io import BytesIO

from docx import Document

from edusci.reporting.docx import build_report_docx
from tests.test_report_v2_contracts import valid_report_payload


def test_v2_docx_has_a4_academic_structure_tables_and_appendix() -> None:
    report = valid_report_payload()
    report["appendices"] = {
        "questionnaire": {
            "title": "大学生AI体验调查问卷",
            "items": [
                {"id": "Q1", "dimension": "AI焦虑", "text": "面对AI发展，我感到不安。", "type": "likert5"},
                {"id": "Q2", "dimension": "技术自我效能", "text": "我能判断AI输出是否可靠。", "type": "likert5"},
            ],
        },
        "analysis_details": {"reliability": {"AI焦虑": {"cronbach_alpha": 0.82}}},
    }
    review = {"overall": "WARN", "checks": {"construct_alignment": {"status": "WARN", "message": "缺少整体心理健康指标"}}}

    content = build_report_docx(report, review, "final")
    document = Document(BytesIO(content))

    section = document.sections[0]
    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.page_height.cm, 1) == 29.7
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for heading in (
        "摘要",
        "待研究问题",
        "解决思路与科学假设",
        "必要的技术手段",
        "数据集",
        "方法论",
        "实验设计",
        "实验结果",
        "参考文献",
        "附录A：调查问卷",
    ):
        assert heading in text
    assert len(document.tables) >= 5
    assert "{'" not in text
    assert "AI焦虑" in " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "科学假设与研究计划" in section.header.paragraphs[0].text
    assert "WARN" in section.footer.paragraphs[0].text
