from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CONTENT_WIDTH_DXA = 9026
NAVY = "183B56"
BLUE = "2E6078"
LIGHT_BLUE = "EAF1F5"
LIGHT_GRAY = "F3F5F6"
GOLD = "A67822"
MUTED = "5E6A71"


def _font(run, name: str, size: float, *, bold: bool = False, color: str = "222222"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def _set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _add_table(document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        _shade(cell, LIGHT_BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(paragraph.runs[0], "Microsoft YaHei", 9, bold=True, color=NAVY)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            _font(paragraph.runs[0], "SimSun", 9)
    _set_table_geometry(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _heading(document, text: str, level: int = 1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _body(document, text: str, *, bold_lead: str | None = None):
    paragraph = document.add_paragraph()
    if bold_lead:
        _font(paragraph.add_run(bold_lead), "SimSun", 10.5, bold=True)
    _font(paragraph.add_run(text), "SimSun", 10.5)
    return paragraph


def _bullet(document, text: str, numbered: bool = False):
    paragraph = document.add_paragraph(style="List Number" if numbered else "List Bullet")
    _font(paragraph.add_run(text), "SimSun", 10.5)
    return paragraph


def _citation_suffix(evidence_ids: list[str], reference_index: dict[str, int]) -> str:
    numbers = sorted({reference_index[eid] for eid in evidence_ids if eid in reference_index})
    return "" if not numbers else " " + "".join(f"[{number}]" for number in numbers)


def _style_document(document: Document, report: dict, review: dict) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    normal = document.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Pt(21)

    for name, size, before, after, color in (
        ("Heading 1", 15, 16, 8, NAVY),
        ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 11, 10, 4, BLUE),
    ):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(header.add_run("科学假设与研究计划 · V2"), "Microsoft YaHei", 8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    overall = review.get("overall", "尚未复审")
    _font(footer.add_run(f"V2 · {overall}  |  第 "), "Microsoft YaHei", 8.5, color=MUTED)
    _add_page_field(footer)
    _font(footer.add_run(" 页"), "Microsoft YaHei", 8.5, color=MUTED)

    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def _build_v2(report: dict, review: dict, mode: str) -> bytes:
    document = Document()
    _style_document(document, report, review)
    references = report.get("references", [])
    reference_index = {item["evidence_id"]: index + 1 for index, item in enumerate(references)}
    overall = review.get("overall", "尚未复审")
    edition = "最终版" if overall == "PASS" else "条件版" if overall == "WARN" else "草稿版"

    cover_spacer = document.add_paragraph()
    cover_spacer.paragraph_format.space_after = Pt(72)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(kicker.add_run("EDUSCI · COMPETITION RESEARCH REPORT"), "Microsoft YaHei", 9, bold=True, color=GOLD)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(18)
    _font(title.add_run(report["paper_title"]), "Microsoft YaHei", 22, bold=True, color=NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(subtitle.add_run("《科学假设与研究计划》"), "SimSun", 14, bold=True, color=BLUE)
    document.add_paragraph().paragraph_format.space_after = Pt(52)
    metadata = [
        ["报告版本", f"V{report.get('schema_version', 2)} · {edition}"],
        ["研究路径", str(report.get("route", "教育研究验证"))],
        ["结果性质", report["results"]["kind"]],
        ["复审状态", overall],
    ]
    _add_table(document, ["项目", "内容"], metadata, [2200, 6826])
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(note.add_run("本报告严格区分真实数据、官方公开数据、模拟输入与拟采集Target数据。"), "SimSun", 9, color=MUTED)
    document.add_page_break()

    _heading(document, "目录")
    for item in (
        "摘要与关键词",
        "1 待研究问题",
        "2 解决思路与科学假设",
        "3 必要的技术手段",
        "4 数据集",
        "5 方法论",
        "6 实验设计",
        "7 实验结果",
        "8 局限、伦理与复审",
        "参考文献",
        "附录",
    ):
        _body(document, item)
    _heading(document, "摘要")
    _body(document, report["abstract"])
    _body(document, "；".join(report["keywords"]), bold_lead="关键词：")

    _heading(document, "1 待研究问题")
    problem = report["problem_statement"]
    suffix = _citation_suffix(problem.get("evidence_ids", []), reference_index)
    _body(document, problem["current_limitation"] + suffix, bold_lead="当前局限：")
    _body(document, problem["knowledge_gap"] + suffix, bold_lead="知识缺口：")
    _body(document, problem["research_question"], bold_lead="研究问题：")

    _heading(document, "2 解决思路与科学假设")
    rationale = report["rationale"]
    _body(document, rationale["innovation"] + _citation_suffix(rationale.get("evidence_ids", []), reference_index), bold_lead="创新点：")
    _heading(document, "2.1 推导链条", 2)
    for step in rationale["reasoning_chain"]:
        _bullet(document, step, numbered=True)
    _heading(document, "2.2 可证伪假设", 2)
    hypothesis_rows = [
        [item["id"], item["null_hypothesis"], item["alternative_hypothesis"], item["falsification_criterion"]]
        for item in report["hypotheses"]
    ]
    _add_table(document, ["编号", "零假设 H0", "备择假设 H1", "证伪标准"], hypothesis_rows, [700, 2450, 2450, 3426])

    document.add_page_break()
    _heading(document, "3 必要的技术手段")
    technical_rows = [
        [item["purpose"], item["method"], " / ".join(item["stack"]), item["execution_status"], item["rationale"]]
        for item in report["technical_details"]
    ]
    _add_table(document, ["目的", "方法", "技术栈", "状态", "选择理由"], technical_rows, [1300, 1900, 1500, 1200, 3126])

    _heading(document, "4 数据集")
    datasets = report["datasets"]
    _heading(document, "4.1 Source：真实历史或公开数据", 2)
    if datasets["source"]:
        source_rows = [[item["name"], item["origin_type"], str(item.get("rows") or "-"), item.get("provenance_url", ""), item.get("license_name", "")] for item in datasets["source"]]
        _add_table(document, ["数据集", "类型", "样本量", "来源", "许可"], source_rows, [1800, 1400, 900, 3200, 1726])
    else:
        _body(document, "当前没有满足来源确认与许可要求的真实Source数据。")
    _heading(document, "4.2 模拟输入", 2)
    if datasets["simulation_input"]:
        simulation_rows = [[item["name"], str(item.get("rows") or "-"), item["role"], "；".join(item.get("limitations", []))] for item in datasets["simulation_input"]]
        _add_table(document, ["数据集", "样本量", "用途", "限制"], simulation_rows, [2100, 1000, 2500, 3426])
    else:
        _body(document, "未使用模拟数据。")
    _heading(document, "4.3 Target：拟采集验证数据", 2)
    target = datasets["target"]
    target_rows = [
        ["目标人群", target["population"]],
        ["特征", "；".join(target["features"])],
        ["结果变量", target["label"]],
        ["最低样本量", str(target["minimum_sample_size"])],
        ["采集周期", target["collection_period"]],
        ["格式", target["format"]],
        ["伦理要求", "；".join(target["ethics"])],
    ]
    _add_table(document, ["字段", "要求"], target_rows, [2000, 7026])

    _heading(document, "5 方法论")
    method_rows = [[str(item["step"]), item["name"], item["input"], item["procedure"], item["output"]] for item in report["methods"]]
    _add_table(document, ["步骤", "环节", "输入", "实施过程", "输出"], method_rows, [650, 1300, 1600, 3476, 2000])

    _heading(document, "6 实验设计")
    experiments = report["experiments"]
    _heading(document, "6.1 Baselines", 2)
    _add_table(document, ["基线", "定义", "用途"], [[item["name"], item["description"], item["purpose"]] for item in experiments["baselines"]], [1800, 3500, 3726])
    _heading(document, "6.2 Metrics", 2)
    _add_table(document, ["指标", "定义", "判定标准"], [[item["name"], item["definition"], item["success_criterion"]] for item in experiments["metrics"]], [1800, 3200, 4026])
    _body(document, experiments["validation_design"], bold_lead="验证设计：")
    for check in experiments.get("robustness_checks", []):
        _bullet(document, check)

    document.add_page_break()
    _heading(document, "7 实验结果")
    results = report["results"]
    _body(document, results["status"], bold_lead="结果状态：")
    _body(document, results["feasibility_conclusion"], bold_lead="可行性结论：")
    _body(document, str(results.get("sample_size") or "未采集"), bold_lead="有效样本量：")
    findings = results.get("statistical_findings", [])
    if findings:
        finding_rows = []
        for item in findings:
            interval = item.get("confidence_interval_95") or []
            ci_text = "-" if not interval else f"[{interval[0]}, {interval[1]}]"
            finding_rows.append([
                item["method"], " / ".join(item["variables"]), str(item.get("estimate")),
                str(item.get("p_value")), ci_text, item["interpretation"],
            ])
        _add_table(document, ["方法", "变量", "估计值", "p值", "95%CI", "解释"], finding_rows, [1300, 1900, 1000, 800, 1500, 2526])
    else:
        _body(document, "当前没有可报告的实证统计量；以下内容仅为验证公式与决策标准。")
    _heading(document, "7.1 验证公式", 2)
    for formula in results.get("formulas", []):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(paragraph.add_run(formula), "Cambria Math", 11, color=NAVY)
    for limitation in results.get("limitations", []):
        _bullet(document, limitation)

    _heading(document, "8 局限、伦理与复审")
    limits = report["limitations_ethics"]
    _body(document, limits["causal_boundary"], bold_lead="因果边界：")
    for label, values in (("样本限制", limits["sample_limitations"]), ("隐私", limits["privacy"]), ("知情同意", limits["consent"])):
        _body(document, "；".join(values) or "无", bold_lead=f"{label}：")
    review_rows = [[name, item.get("status", ""), item.get("message", "")] for name, item in review.get("checks", {}).items()]
    if review_rows:
        _add_table(document, ["复审项", "状态", "说明"], review_rows, [1900, 1000, 6126])

    if report.get("schema_version", 2) >= 3:
        _heading(document, "9 证据矩阵与结论")
        methodology = report.get("research_methodology", {})
        _body(
            document,
            (
                f"完成 {methodology.get('iterations', 0)} 轮证据检索，"
                f"解析 {methodology.get('fulltexts_parsed', 0)} 篇开放全文。"
            ),
        )
        conclusions = report.get("conclusions", [])
        if conclusions:
            _add_table(
                document,
                ["观点", "置信度", "证据ID"],
                [
                    [
                        item["statement"],
                        item["confidence"],
                        "、".join(item.get("evidence_ids", [])),
                    ]
                    for item in conclusions
                ],
                [4800, 1300, 2926],
            )
        evidence = report.get("evidence_ledger", [])
        if evidence:
            _add_table(
                document,
                ["来源", "立场", "定位", "证据片段", "对应观点"],
                [
                    [
                        item.get("document_title", ""),
                        item.get("stance", ""),
                        item.get("locator", ""),
                        item.get("excerpt", ""),
                        next(
                            (
                                conclusion["statement"]
                                for conclusion in conclusions
                                if conclusion["claim_id"] == item.get("claim_id")
                            ),
                            item.get("claim_id", ""),
                        ),
                    ]
                    for item in evidence
                ],
                [1600, 900, 1000, 3000, 2526],
            )
        for conflict in report.get("conflicting_evidence", []):
            _bullet(document, f"冲突证据：{conflict.get('statement', '')}")

    document.add_page_break()
    _heading(document, "参考文献")
    if references:
        for index, reference in enumerate(references, start=1):
            _body(document, reference["formatted"], bold_lead=f"[{index}] ")
    else:
        _body(document, "当前没有可核验参考文献，报告不得作为最终版发布。")

    appendices = report.get("appendices", {})
    questionnaire = appendices.get("questionnaire") or {}
    if questionnaire:
        document.add_page_break()
        _heading(document, "附录A：调查问卷")
        _body(document, questionnaire.get("title", "调查问卷"))
        question_rows = [[item.get("id", ""), item.get("dimension", ""), item.get("text", ""), item.get("type", "")] for item in questionnaire.get("items", [])]
        _add_table(document, ["题号", "维度", "题目", "作答形式"], question_rows, [700, 1500, 5226, 1600])
    details = appendices.get("analysis_details") or {}
    if details:
        _heading(document, "附录B：分析审计信息")
        reliability_rows = []
        for scale, values in details.get("reliability", {}).items():
            reliability_rows.append([scale, str(values.get("cronbach_alpha")), str(values.get("spearman_brown", "-"))])
        if reliability_rows:
            _add_table(document, ["量表", "Cronbach's α", "Spearman-Brown"], reliability_rows, [4000, 2513, 2513])

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _build_legacy(report: dict, review: dict, mode: str) -> bytes:
    document = Document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(title.add_run(report.get("title", "教育学科学假设与研究计划")), "Microsoft YaHei", 20, bold=True)
    for key, label in (("abstract", "摘要"), ("problem", "待研究问题"), ("rationale", "解决思路"), ("methods", "方法论"), ("results", "实验结果")):
        _heading(document, label)
        _body(document, str(report.get(key, "")))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def build_report_docx(report: dict, review: dict, mode: str) -> bytes:
    if report.get("schema_version", 0) >= 2:
        return _build_v2(report, review, mode)
    return _build_legacy(report, review, mode)
